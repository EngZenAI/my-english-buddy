import csv
import io

from fastapi import APIRouter, File, UploadFile

from backend.db.dependencies import SessionDep
from backend.db.repositories import (
    add_label,
    bulk_delete_words,
    bulk_update_words,
    delete_word,
    existing_words_lower,
    get_all_words,
    insert_words,
    is_word_saved,
    reorder_words,
    save_word,
    update_word,
)
from backend.exceptions import FILE_IMPORT_ERRORS
from backend.routers.common import CurrentUserDep
from backend.schemas.wordbook import (
    BulkUpdateIn,
    IdsIn,
    ImportCommitIn,
    SaveWordIn,
    UpdateWordIn,
)
from backend.wordbook.starter_packs import get_starter_pack, list_starter_packs

router = APIRouter(tags=["wordbook"])


@router.get(
    "/words/saved",
    summary="단어 저장 여부 확인",
    description="검색 결과에 표시할 수 있도록 현재 사용자의 단어장에 이미 저장된 단어인지 확인합니다.",
)
async def word_saved(word: str = "", *, session: SessionDep, _user: CurrentUserDep):
    return {"saved": await is_word_saved(session, _user["id"], word)}


@router.get(
    "/words",
    summary="단어장 목록 조회",
    description=(
        "현재 사용자의 저장 단어를 정렬 순서대로 반환합니다. "
        "tag 쿼리가 있으면 해당 태그의 단어만 조회합니다."
    ),
)
async def list_words(tag: str = "", *, session: SessionDep, _user: CurrentUserDep):
    return {"words": await get_all_words(session, _user["id"], tag or None)}


@router.post(
    "/words",
    summary="단어 저장",
    description=(
        "검색 화면, 롤플레잉, 뉴스 리딩 등에서 선택한 단어를 현재 사용자의 단어장에 저장합니다. "
        "태그가 비어 있으면 '미지정'으로 저장합니다."
    ),
)
async def create_word(payload: SaveWordIn, session: SessionDep, _user: CurrentUserDep):
    word = payload.word.strip()
    if not word:
        return {"ok": False, "message": "단어가 비어 있습니다."}
    final_def = payload.slang_def.strip() or payload.english_def
    tag = payload.tag.strip() or "미지정"
    message = await save_word(
        session,
        _user["id"],
        word.lower(),
        payload.korean,
        payload.korean_detail,
        final_def,
        payload.example,
        tag,
    )
    return {"ok": True, "message": message, "saved": True}


@router.patch(
    "/words/{word_id}",
    summary="단어 단건 수정",
    description="단어장 한 항목의 상세 뜻, 영어 뜻, 예문, 태그, 복습일을 수정합니다.",
)
async def edit_word(word_id: int, payload: UpdateWordIn, session: SessionDep, _user: CurrentUserDep):
    ok = await update_word(
        session,
        _user["id"],
        word_id,
        payload.korean_detail,
        payload.english_def,
        payload.example,
        payload.tag.strip(),
        payload.next_review.strip(),
    )
    if not ok:
        return {"ok": False, "message": "수정할 단어를 찾을 수 없어요."}
    return {"ok": True, "message": "✏️ 수정했어요!"}


@router.delete(
    "/words/{word_id}",
    summary="단어 단건 삭제",
    description="현재 사용자의 단어장에서 지정한 단어 항목을 삭제합니다.",
)
async def remove_word(word_id: int, session: SessionDep, _user: CurrentUserDep):
    ok = await delete_word(session, _user["id"], word_id)
    if not ok:
        return {"ok": False, "message": "삭제할 단어를 찾을 수 없어요."}
    return {"ok": True, "message": "🗑️ 삭제했어요!"}


@router.post(
    "/words/bulk-update",
    summary="단어장 일괄 수정",
    description=(
        "단어장 편집 모드에서 여러 단어의 본문, 뜻, 예문, 태그, 복습일을 한 번에 저장합니다. "
        "영어 단어 변경 시 사용자 단어장 내 중복 단어는 건너뜁니다."
    ),
)
async def bulk_update(payload: BulkUpdateIn, session: SessionDep, _user: CurrentUserDep):
    items = [it.model_dump() for it in payload.items]
    res = await bulk_update_words(session, _user["id"], items)
    updated = res.get("updated", 0)
    conflicts = res.get("conflicts", [])
    if not res.get("ok", True):
        return {
            "ok": False,
            "updated": 0,
            "skipped": res.get("skipped", 0),
            "conflicts": conflicts,
            "message": res.get("message", "중복으로 저장하지 못했어요."),
        }
    message = f"💾 {updated}개 저장했어요!"
    if conflicts:
        message += (
            f" ({len(conflicts)}개는 이미 있는 단어와 겹쳐 건너뜀: "
            + ", ".join(conflicts)
            + ")"
        )
    return {
        "ok": True,
        "updated": updated,
        "skipped": res.get("skipped", 0),
        "conflicts": conflicts,
        "message": message,
    }


@router.post(
    "/words/bulk-delete",
    summary="단어장 다중 삭제",
    description="선택한 단어 ID 목록을 현재 사용자의 단어장에서 한 번에 삭제합니다.",
)
async def bulk_delete(payload: IdsIn, session: SessionDep, _user: CurrentUserDep):
    deleted = await bulk_delete_words(session, _user["id"], payload.ids)
    return {"ok": True, "deleted": deleted, "message": f"🗑️ {deleted}개 삭제했어요!"}


@router.post(
    "/words/reorder",
    summary="단어장 순서 저장",
    description="드래그 정렬 결과로 전달된 단어 ID 순서를 sort_order에 반영합니다.",
)
async def reorder(payload: IdsIn, session: SessionDep, _user: CurrentUserDep):
    updated = await reorder_words(session, _user["id"], payload.ids)
    return {"ok": True, "updated": updated}


def _decode_bytes(content: bytes) -> str:
    """구글 내보내기는 UTF-8, Excel 재저장본은 CP949일 수 있어 차례로 시도."""
    for enc in ("utf-8-sig", "cp949", "utf-8"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


_LANG_EN = {"영어", "english", "en", "eng"}
_LANG_KO = {"한국어", "korean", "ko", "kor"}
_ALL_LANG = _LANG_EN | _LANG_KO


def _looks_like_header(a: str, b: str) -> bool:
    s = (a + " " + b).lower()
    keys = ("영어", "한국어", "english", "word", "korean", "뜻", "translation", "단어")
    return any(k in s for k in keys)


def _row_to_pair(cells) -> tuple[str, str] | None:
    """한 행을 (영어단어, 한국어뜻)으로 변환.
    구글 번역 내보내기는 4열 [소스언어, 타깃언어, 소스텍스트, 타깃텍스트]이고
    행마다 영→한/한→영 방향이 다르므로 언어 라벨을 보고 영어 쪽을 단어로 잡는다.
    그 외에는 단순 2열 [영어, 한국어]로 처리."""
    cells = [("" if c is None else str(c)).strip() for c in cells]
    if (
        len(cells) >= 4
        and cells[0].lower() in _ALL_LANG
        and cells[1].lower() in _ALL_LANG
    ):
        src_lang, tgt_lang, src_txt, tgt_txt = (
            cells[0].lower(),
            cells[1].lower(),
            cells[2],
            cells[3],
        )
        if src_lang in _LANG_EN:
            return src_txt, tgt_txt
        if tgt_lang in _LANG_EN:
            return tgt_txt, src_txt
        return src_txt, tgt_txt
    if len(cells) >= 2:
        return cells[0], cells[1]
    return None


def _rows_from_records(records) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    for i, r in enumerate(records):
        pair = _row_to_pair(r)
        if not pair:
            continue
        word, korean = pair
        if not word:
            continue
        if i == 0 and len(list(r)) < 3 and _looks_like_header(word, korean):
            continue
        rows.append((word, korean))
    return rows


def _parse_rows_csv(content: bytes) -> list[tuple[str, str]]:
    text = _decode_bytes(content)
    return _rows_from_records(csv.reader(io.StringIO(text)))


def _parse_rows_xlsx(content: bytes) -> list[tuple[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    return _rows_from_records(ws.iter_rows(values_only=True))


def _parse_upload(file: UploadFile, content: bytes) -> list[tuple[str, str]]:
    name = (file.filename or "").lower()
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        return _parse_rows_xlsx(content)
    return _parse_rows_csv(content)


@router.post(
    "/words/import/preview",
    summary="단어 가져오기 미리보기",
    description=(
        "CSV/XLSX 파일을 파싱해서 저장 후보 행을 반환합니다. "
        "이 단계에서는 DB에 저장하지 않고, 이미 보유한 단어 여부만 표시합니다."
    ),
)
async def import_preview(
    *,
    file: UploadFile = File(...),
    session: SessionDep,
    _user: CurrentUserDep,
):
    """파일을 파싱만 해서 미리보기 행을 돌려준다(저장하지 않음).
    각 행에 이미 보유 중인 단어인지(dup) 표시한다."""
    content = await file.read()
    try:
        rows = _parse_upload(file, content)
    except FILE_IMPORT_ERRORS:
        return {"ok": False, "message": "파일을 읽지 못했어요. CSV 또는 XLSX인지, 첫 두 열이 영어/한국어인지 확인해주세요."}

    if not rows:
        return {"ok": False, "message": "가져올 단어가 없어요. 첫 두 열이 '영어 | 한국어' 형식인지 확인해주세요."}

    existing = await existing_words_lower(session, _user["id"])
    out = [
        {"word": w, "korean": k, "dup": w.lower() in existing}
        for w, k in rows
    ]
    return {"ok": True, "rows": out, "count": len(out)}


@router.post(
    "/words/import/commit",
    summary="단어 가져오기 적용",
    description="미리보기 모달에서 검토·수정한 행을 실제 단어장에 저장합니다.",
)
async def import_commit(payload: ImportCommitIn, session: SessionDep, _user: CurrentUserDep):
    """미리보기에서 검토·편집한 행들을 실제로 저장한다."""
    items = [it.model_dump() for it in payload.items]
    if not items:
        return {"ok": False, "message": "적용할 단어가 없어요."}
    result = await insert_words(session, _user["id"], items, overwrite=payload.overwrite)
    parts = [f"📥 {result['added']}개 추가"]
    if result.get("updated"):
        parts.append(f"{result['updated']}개 덮어씀")
    if result.get("skipped"):
        parts.append(f"{result['skipped']}개는 이미 있어 건너뜀")
    return {
        "ok": True,
        "message": ", ".join(parts),
        **result,
    }


@router.get("/words/starter-packs")
async def starter_packs(session: SessionDep, _user: CurrentUserDep):
    """기본 제공(예시) 단어장 목록. 각 팩이 이미 몇 개 담겨 있는지도 함께 알려준다."""
    existing = await existing_words_lower(session, _user["id"])
    out = []
    for pack in list_starter_packs():
        words = pack["words"]
        already = sum(1 for w in words if w["word"].lower() in existing)
        out.append(
            {
                "id": pack["id"],
                "title": pack["title"],
                "tag": pack["tag"],
                "description": pack["description"],
                "count": len(words),
                "already": already,
            }
        )
    return {"packs": out}


@router.post("/words/starter-packs/{pack_id}/import")
async def import_starter_pack(pack_id: str, session: SessionDep, _user: CurrentUserDep):
    """기본 제공 단어장을 사용자 단어장에 담는다.
    팩의 태그를 없으면 만들고, 이미 있는 단어는 건너뛴다(중복 방지).
    삭제 후 다시 눌러도 빠진 단어만 다시 채워진다."""
    pack = get_starter_pack(pack_id)
    if not pack:
        return {"ok": False, "message": "해당 단어장을 찾을 수 없어요."}
    tag = pack["tag"]
    await add_label(session, _user["id"], tag)
    items = [
        {"word": w["word"], "korean": w.get("korean", ""), "tag": tag}
        for w in pack["words"]
    ]
    result = await insert_words(session, _user["id"], items, overwrite=False)
    added = result.get("added", 0)
    skipped = result.get("skipped", 0)
    if added:
        message = f"'{pack['title']}' 단어장에서 {added}개를 담았어요!"
        if skipped:
            message += f" ({skipped}개는 이미 있어 건너뜀)"
    else:
        message = f"이미 '{pack['title']}' 단어장을 모두 담았어요."
    return {"ok": True, "message": message, "tag": tag, **result}
